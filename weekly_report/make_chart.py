import pandas as pd
from loguru import logger

from weekly_report.utils.utils import max_used_col


def make_chart(writer, df, index_vars, sheet_name):
    """ pivot sincere data by day/master then chart it """

    from openpyxl.chart import (
        LineChart,
        Reference,
    )

    # df_pt = pd.pivot_table(df, columns='created_at', index=index_vars,
    #                        values=['addresses_count'], margins=True, dropna=True, aggfunc=np.sum)
    df_pt = pd.pivot_table(df, columns='created_at', index=index_vars,
                           values=['addresses_count'], margins=True, dropna=True, aggfunc='sum')
    df_pt = df_pt.sort_index(axis=1, ascending=False)
    df_pt = df_pt.drop(df_pt.columns[0], axis=1)  # all column from margins=True (need all row)
    df_pt.to_excel(writer, sheet_name=sheet_name, startrow=6)
    ws = writer.sheets[sheet_name]
    # mycell = ws['C12']
    # ws.freeze_panes = mycell
    # ws.freeze_panes(freeze_row, freeze_col)

    # for r in dataframe_to_rows(df_pt, index=True, header=True):
    #     ws.append(r)

    c1 = LineChart()
    c1.title = "Daily Sincere Writer Address Requests"
    c1.style = 13
    c1.y_axis.title = 'Addresses'
    c1.x_axis.title = 'Date'

    # chart.y_axis.crossAx = 500
    # chart.x_axis = DateAxis(crossAx=100)
    # chart.x_axis.number_format ='yyyy/mm/dd'
    # chart.x_axis.majorTimeUnit = "days"

    used_col = max_used_col(ws, 8)

    data = Reference(ws, min_col=1, max_col=used_col, min_row=10, max_row=ws.max_row)
    cats = Reference(ws, min_col=2, max_col=used_col, min_row=8, max_row=8)
    c1.add_data(data, from_rows=True, titles_from_data=True)
    c1.x_axis.number_format = 'mm/dd'

    c1.height = 10 * 2  # default is 7.5
    c1.width = 17 * 2  # default is 15
    c1.legend.position = 'b'  # 'tr', 'b', 't', 'l', 'r'

    c1.set_categories(cats)
    ws.add_chart(c1, "A1")
    # ws.add_chart(c1, "A1")
    # wb.save("SampleChart.xlsx")
    # writer.save()
    logger.debug("leaving make_chart")
